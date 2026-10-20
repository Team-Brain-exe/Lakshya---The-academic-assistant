import io
import unittest

from openpyxl import load_workbook

from app.engine.co_math import (
    COConfig,
    CountedScore,
    StudentInfo,
    build_co_sheet,
    class_summary,
)
from app.engine.grading import COMarks
from app.export.sheet_xlsx import build_sheet_workbook

CONFIG = [COConfig(1, 60, 60), COConfig(2, 60, 60), COConfig(3, 60, 60)]
THRESHOLDS = {c.number: c.threshold_pct for c in CONFIG}


def counted(usn, quiz, module, **by_co):
    marks = {int(k[2:]): COMarks(int(k[2:]), ob, mx) for k, (ob, mx) in by_co.items()}
    return CountedScore(usn, quiz, module, marks)


def make_sheet(roster=None):
    # S1: CO1 2/3 = 66.67, CO2 1/2 = 50, total 3/5 = 60
    # S2: CO1 0/3 = 0,     CO2 2/2 = 100, total 2/5 = 40
    # S3: no attempts. CO3 has no questions at all.
    data = [
        counted("S1", "Q1", "M1", co1=(2, 3), co2=(1, 2)),
        counted("S2", "Q1", "M1", co1=(0, 3), co2=(2, 2)),
    ]
    roster = roster or [StudentInfo("S1", "Asha"), StudentInfo("S2", "Bala"), StudentInfo("S3", "Cara")]
    return build_co_sheet(data, roster, CONFIG)


def open_xlsx(data):
    return load_workbook(io.BytesIO(data))


def rows(ws):
    return [tuple(r) for r in ws.iter_rows(values_only=True)]


class SheetExportTests(unittest.TestCase):
    def test_header_rows_and_average_match_the_sheet(self):
        wb = open_xlsx(build_sheet_workbook(make_sheet()))
        self.assertEqual(
            rows(wb["CO sheet"]),
            [
                ("USN", "Name", "CO1 (%)", "CO2 (%)", "CO3 (%)", "Total (%)"),
                ("S1", "Asha", 66.67, 50, None, 60),
                ("S2", "Bala", 0, 100, None, 40),
                ("S3", "Cara", None, None, None, None),
                ("Average", None, 33.33, 75, None, None),
            ],
        )

    def test_no_data_is_blank_never_zero(self):
        ws = open_xlsx(build_sheet_workbook(make_sheet()))["CO sheet"]
        self.assertIsNone(ws["E2"].value)  # CO3 has no questions
        self.assertIsNone(ws["C4"].value)  # S3 has no attempts
        self.assertEqual(ws["C3"].value, 0)  # a real zero stays a zero

    def test_every_exported_row_equals_the_on_screen_row(self):
        sheet = make_sheet()
        ws = open_xlsx(build_sheet_workbook(sheet))["CO sheet"]
        for i, row in enumerate(sheet.rows, start=2):
            self.assertEqual(ws.cell(row=i, column=1).value, row.usn)
            for j, co in enumerate(sheet.co_numbers, start=3):
                shown = None if row.co_scores[co] is None else round(row.co_scores[co], 2)
                self.assertEqual(ws.cell(row=i, column=j).value, shown)
            total = None if row.total_pct is None else round(row.total_pct, 2)
            self.assertEqual(ws.cell(row=i, column=6).value, total)

    def test_section_column_appears_only_when_sections_exist(self):
        roster = [StudentInfo("S1", "Asha", "A"), StudentInfo("S2", "Bala", "B")]
        ws = open_xlsx(build_sheet_workbook(make_sheet(roster)))["CO sheet"]
        self.assertEqual(
            rows(ws)[0], ("USN", "Name", "Section", "CO1 (%)", "CO2 (%)", "CO3 (%)", "Total (%)")
        )
        self.assertEqual(rows(ws)[1][:3], ("S1", "Asha", "A"))

    def test_scores_below_the_threshold_are_shaded(self):
        ws = open_xlsx(build_sheet_workbook(make_sheet(), thresholds=THRESHOLDS))["CO sheet"]
        self.assertEqual(ws["C3"].fill.fill_type, "solid")  # S2 CO1 = 0
        self.assertEqual(ws["D2"].fill.fill_type, "solid")  # S1 CO2 = 50
        self.assertIsNone(ws["C2"].fill.fill_type)  # S1 CO1 = 66.67
        self.assertIsNone(ws["E3"].fill.fill_type)  # no data is not "below"

    def test_a_score_exactly_at_the_threshold_is_not_shaded(self):
        sheet = build_co_sheet(
            [counted("S1", "Q1", "M1", co1=(3, 5))], [StudentInfo("S1")], [COConfig(1, 60, 60)]
        )
        ws = open_xlsx(build_sheet_workbook(sheet, thresholds={1: 60}))["CO sheet"]
        self.assertEqual(ws["C2"].value, 60)
        self.assertIsNone(ws["C2"].fill.fill_type)

    def test_no_shading_without_thresholds(self):
        ws = open_xlsx(build_sheet_workbook(make_sheet()))["CO sheet"]
        self.assertIsNone(ws["C3"].fill.fill_type)

    def test_text_that_looks_like_a_formula_stays_text(self):
        sheet = make_sheet([StudentInfo("=1+1", "+cmd")])
        ws = open_xlsx(build_sheet_workbook(sheet))["CO sheet"]
        self.assertEqual(ws["A2"].value, "=1+1")
        self.assertEqual(ws["A2"].data_type, "s")
        self.assertEqual(ws["B2"].value, "+cmd")

    def test_empty_roster_gives_a_header_and_an_average_row(self):
        sheet = build_co_sheet([], [], CONFIG)
        ws = open_xlsx(build_sheet_workbook(sheet))["CO sheet"]
        self.assertEqual(
            rows(ws),
            [
                ("USN", "Name", "CO1 (%)", "CO2 (%)", "CO3 (%)", "Total (%)"),
                ("Average", None, None, None, None, None),
            ],
        )

    def test_sheet_title_is_made_safe_for_excel(self):
        wb = open_xlsx(build_sheet_workbook(make_sheet(), title="Data Structures: M1/M2"))
        self.assertEqual(wb.sheetnames, ["Data Structures- M1-M2"])
        long = open_xlsx(build_sheet_workbook(make_sheet(), title="x" * 60))
        self.assertEqual(len(long.sheetnames[0]), 31)
        blank = open_xlsx(build_sheet_workbook(make_sheet(), title="   "))
        self.assertEqual(blank.sheetnames, ["CO sheet"])

    def test_the_file_is_a_real_xlsx(self):
        data = build_sheet_workbook(make_sheet())
        self.assertEqual(data[:2], b"PK")  # xlsx is a zip file


class SummaryExportTests(unittest.TestCase):
    def test_summary_sheet_is_added_when_given(self):
        sheet = make_sheet()
        summary = class_summary(sheet, CONFIG)
        wb = open_xlsx(build_sheet_workbook(sheet, summary=summary))
        self.assertEqual(wb.sheetnames, ["CO sheet", "Class summary"])
        got = rows(wb["Class summary"])
        self.assertEqual(
            got[0],
            (
                "CO",
                "Threshold (%)",
                "Target (%)",
                "Enrolled",
                "Attempted",
                "Average (%)",
                "At or above threshold",
                "Below threshold",
                "Attainment (%)",
                "Target met",
            ),
        )
        # CO1: S1 66.67 reaches 60, S2 0 does not; 1 of 3 enrolled = 33.33, target 60 not met
        self.assertEqual(got[1], ("CO1", 60, 60, 3, 2, 33.33, 1, 1, 33.33, "No"))
        # CO2: S1 50 below, S2 100 above
        self.assertEqual(got[2], ("CO2", 60, 60, 3, 2, 75, 1, 1, 33.33, "No"))
        # CO3: nobody has data
        self.assertEqual(got[3], ("CO3", 60, 60, 3, 0, None, 0, 0, 0, "No data"))

    def test_summary_sheet_is_absent_by_default(self):
        wb = open_xlsx(build_sheet_workbook(make_sheet()))
        self.assertEqual(wb.sheetnames, ["CO sheet"])


if __name__ == "__main__":
    unittest.main()
