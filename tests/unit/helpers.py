"""Shared test helpers."""

import io

from pptx import Presentation
from pptx.util import Inches


def build_deck() -> bytes:
    """A small deck: slides 1, 2 and 4 have text; slide 3 is empty."""
    prs = Presentation()

    # Slide 1: title and content layout, with speaker notes
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Hash tables"
    s1.placeholders[1].text = "Keys map to buckets\nCollisions need a strategy"
    s1.notes_slide.notes_text_frame.text = "Mention load factor"

    # Slide 2: blank slide with a text box, a table and a group
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    s2.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "Chaining vs probing"
    table = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Method"
    table.cell(0, 1).text = "Cost"
    table.cell(1, 0).text = "Chaining"
    table.cell(1, 1).text = "O(1) average"
    group = s2.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(4), Inches(3), Inches(1)).text_frame.text = "Grouped note"

    # Slide 3: blank, no text (stands in for a picture-only slide)
    prs.slides.add_slide(prs.slide_layouts[6])

    # Slide 4: title only
    s4 = prs.slides.add_slide(prs.slide_layouts[5])
    s4.shapes.title.text = "Summary"

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build_picture_only_deck() -> bytes:
    """Two slides, no text anywhere."""
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.slides.add_slide(prs.slide_layouts[6])
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
