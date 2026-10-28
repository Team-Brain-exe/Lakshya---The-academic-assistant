"""Integration tests for Arya's routes, run against the real Postgres database.

Covers: permissions, upload status transitions, quiz submission -> CO sheet
consistency (including the `counted` flag), the Excel export, coverage, slide Q&A,
study material and plan, the calendar, and usage logging.

No Token Factory key is needed. AI-backed paths are tested with a fake client that
returns canned text and token counts, which proves the wiring and the logging, not
the quality of any model. Real-credential tests belong in tests/integration/.

Every test creates its own rows through the `ctx` fixture (which removes them) and
never deletes anything it did not create.
"""

from __future__ import annotations

import io
import uuid
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

import openpyxl
import pytest
from pptx import Presentation
from sqlalchemy import text

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
NOT_A_FACULTY = {403}


# ---------------------------------------------------------------------------
# Helpers and fixtures
# ---------------------------------------------------------------------------

@pytest.fixture(autouse=True)
def _no_ai_credentials(monkeypatch):
    """Fallback paths by default: no key, no models."""
    for name in ("TOKEN_FACTORY_API_KEY", "TOKEN_FACTORY_BASE_URL", "MODEL_TAGGING",
                 "MODEL_GENERATION", "MODEL_AGENT", "MODEL_REASONING"):
        monkeypatch.delenv(name, raising=False)


@pytest.fixture
def log_cleanup(db):
    """Remove only the agent_logs rows written by the fake model."""
    yield
    try:
        db.rollback()
        db.execute(text("DELETE FROM agent_logs WHERE model = 'fake-model'"))
        db.commit()
    except Exception:
        db.rollback()


def _tok(ctx, **creds):
    resp = ctx.login(**creds)
    assert resp.status_code == 200, resp.text
    return ctx.headers(resp.json()["access_token"])


def _make_deck(slides: list[tuple[str, str]]) -> bytes:
    prs = Presentation()
    for title, body in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


DECK = _make_deck([
    ("Binary search tree", "A binary search tree keeps smaller keys in the left subtree and larger keys in the right subtree."),
    ("Tree traversal", "Inorder traversal visits left subtree, node, then right subtree and returns keys in sorted order."),
    ("Hash tables", "A hash table maps keys to buckets using a hash function and resolves collisions by chaining."),
])


def _world(ctx, db, *, sections=("A", "B")):
    """Faculty, subject with CO1 and CO2, one module, four approved questions
    (q1,q2 -> CO1 correct option 0; q3,q4 -> CO2 correct option 1), a published
    open quiz and one enrolled student per section."""
    fac = ctx.make_faculty()
    subj = ctx.make_subject(fac["user_id"])
    sid = subj["subject_id"]
    co1 = ctx.make_co(sid, 1, "CO1 trees")
    co2 = ctx.make_co(sid, 2, "CO2 hashing")
    mod = ctx.make_module(sid, 1, "Module 1")
    qs = []
    for i, (co, correct) in enumerate([(co1, 0), (co1, 0), (co2, 1), (co2, 1)]):
        qs.append(ctx.make_question(
            mod["module_id"], co["co_id"], faculty_user_id=fac["user_id"], status="approved",
            stem=f"Question {i + 1}", correct_option=correct, opts=["A", "B", "C", "D"],
        ))
    quiz = ctx.make_quiz(mod["module_id"], sid, faculty_user_id=fac["user_id"],
                         question_ids=[q["question_id"] for q in qs])
    now = datetime.now(timezone.utc)
    db.execute(
        text("UPDATE quizzes SET opens_at = :o, closes_at = :c, published = true, "
             "published_at = now() WHERE id = :id"),
        {"o": now - timedelta(hours=1), "c": now + timedelta(hours=2), "id": quiz["quiz_id"]},
    )
    db.commit()
    students = []
    for sec in sections:
        st = ctx.make_student()
        ctx.make_enrollment(st["student_id"], sid, sec)
        students.append(st)
    return SimpleNamespace(
        fac=fac, sid=sid, co1=co1, co2=co2, mod=mod, qs=qs, quiz=quiz["quiz_id"],
        students=students, fh=_tok(ctx, email=fac["email"], password=fac["password"]),
    )


def _attempt(client, ctx, student, quiz_id, answers_by_index):
    """Start and submit one attempt. answers_by_index: option for q1..q4 (None = skip)."""
    h = _tok(ctx, usn=student["usn"], pin=student["pin"])
    r = client.post(f"/quizzes/{quiz_id}/attempts", headers=h)
    assert r.status_code == 201, r.text
    body = r.json()
    ids = [q["id"] for q in body["questions"]]
    # Questions are returned in quiz order; map by stem to be safe.
    by_stem = {q["id"]: q.get("stem") for q in body["questions"]}
    order = sorted(ids, key=lambda i: by_stem[i])  # "Question 1".."Question 4"
    answers = {qid: opt for qid, opt in zip(order, answers_by_index) if opt is not None}
    sub = client.post(f"/attempts/{body['attempt_id']}/submit", headers=h, json={"answers": answers})
    assert sub.status_code == 200, sub.text
    return body["attempt_id"], sub.json()


ALL_RIGHT = [0, 0, 1, 1]
ALL_WRONG = [3, 3, 3, 3]
CO1_RIGHT_CO2_WRONG = [0, 0, 3, 3]


def _sheet_row(client, w, usn, **params):
    r = client.get(f"/subjects/{w.sid}/co-sheet", headers=w.fh, params=params)
    assert r.status_code == 200, r.text
    return next(row for row in r.json()["rows"] if row["usn"] == usn), r.json()


# ---------------------------------------------------------------------------
# 1. Permissions
# ---------------------------------------------------------------------------

def _faculty_gets(w):
    mid, sid, co = w.mod["module_id"], w.sid, w.co1["co_id"]
    return [
        f"/subjects/{sid}/co-sheet", f"/subjects/{sid}/co-summary", f"/subjects/{sid}/weak-cos",
        f"/subjects/{sid}/coverage", f"/subjects/{sid}/co-sheet/export",
        f"/subjects/{sid}/cos/{co}/drilldown", f"/modules/{mid}/status",
    ]


def test_no_token_is_401_everywhere(client, ctx, db):
    w = _world(ctx, db)
    for path in _faculty_gets(w) + ["/usage/summary", "/events", "/students/me/progress?subject_id=" + w.sid]:
        assert client.get(path).status_code == 401, path
    assert client.post("/qa", json={"subject_id": w.sid, "question": "what is a tree"}).status_code == 401


def test_student_cannot_use_faculty_routes(client, ctx, db):
    w = _world(ctx, db)
    sh = _tok(ctx, usn=w.students[0]["usn"], pin=w.students[0]["pin"])
    for path in _faculty_gets(w) + ["/usage/summary"]:
        r = client.get(path, headers=sh)
        assert r.status_code == 403, (path, r.status_code)
    mid = w.mod["module_id"]
    assert client.post(f"/modules/{mid}/study-material", headers=sh).status_code == 403
    assert client.post(f"/modules/{mid}/upload", headers=sh,
                       files={"file": ("d.pptx", DECK, PPTX_MIME)}).status_code == 403
    r = client.post("/events", headers=sh, json={"subject_id": w.sid, "type": "quiz", "title": "x",
                                                  "starts_at": "2030-01-01T10:00:00Z"})
    assert r.status_code == 403


def test_other_faculty_gets_404_not_data(client, ctx, db):
    w = _world(ctx, db)
    other = ctx.make_faculty()
    oh = _tok(ctx, email=other["email"], password=other["password"])
    for path in _faculty_gets(w):
        r = client.get(path, headers=oh)
        assert r.status_code == 404, (path, r.status_code)
    mid = w.mod["module_id"]
    assert client.post(f"/modules/{mid}/study-material", headers=oh).status_code == 404
    assert client.post(f"/modules/{mid}/upload", headers=oh,
                       files={"file": ("d.pptx", DECK, PPTX_MIME)}).status_code == 404
    r = client.post("/events", headers=oh, json={"subject_id": w.sid, "type": "quiz", "title": "x",
                                                  "starts_at": "2030-01-01T10:00:00Z"})
    assert r.status_code == 404
    # The other faculty member's calendar and usage views never include this subject.
    assert client.get("/events", headers=oh).json() == []
    # The module was not touched by the rejected upload.
    status = db.execute(text("SELECT processing_status FROM modules WHERE id = :m"), {"m": mid}).scalar()
    assert status == "pending"


def test_unenrolled_student_gets_404(client, ctx, db):
    w = _world(ctx, db)
    outsider = ctx.make_student()
    oh = _tok(ctx, usn=outsider["usn"], pin=outsider["pin"])
    assert client.post("/qa", headers=oh, json={"subject_id": w.sid, "question": "what is a tree"}).status_code == 404
    assert client.get("/retrieve", headers=oh, params={"subject_id": w.sid, "q": "tree"}).status_code == 404
    assert client.get("/students/me/progress", headers=oh, params={"subject_id": w.sid}).status_code == 404
    assert client.post("/students/me/study-plan", headers=oh, json={"subject_id": w.sid}).status_code == 404
    assert client.get("/events", headers=oh, params={"subject_id": w.sid}).status_code == 404
    assert client.get(f"/modules/{w.mod['module_id']}/flashcards", headers=oh).status_code == 404


def test_student_sees_only_own_progress(client, ctx, db):
    w = _world(ctx, db)
    a, b = w.students
    _attempt(client, ctx, a, w.quiz, ALL_RIGHT)
    _attempt(client, ctx, b, w.quiz, ALL_WRONG)
    ah = _tok(ctx, usn=a["usn"], pin=a["pin"])

    mine = client.get("/students/me/progress", headers=ah, params={"subject_id": w.sid})
    assert mine.status_code == 200, mine.text
    assert {s["co_number"]: s["percentage"] for s in mine.json()["co_scores"]} == {1: 100.0, 2: 100.0}

    # Another student's USN, in any spelling, is refused.
    for usn in (b["usn"], b["usn"].lower()):
        r = client.get(f"/students/{usn}/progress", headers=ah, params={"subject_id": w.sid})
        assert r.status_code == 403, usn
    # Own USN through the by-USN route works; the faculty member can read both.
    assert client.get(f"/students/{a['usn']}/progress", headers=ah, params={"subject_id": w.sid}).status_code == 200
    rb = client.get(f"/students/{b['usn']}/progress", headers=w.fh, params={"subject_id": w.sid})
    assert rb.status_code == 200
    assert {s["co_number"]: s["percentage"] for s in rb.json()["co_scores"]} == {1: 0.0, 2: 0.0}


def test_malformed_ids_are_404_not_500(client, ctx, db):
    w = _world(ctx, db)
    for path in ("/subjects/not-a-uuid/co-sheet", "/modules/xyz/status",
                 f"/subjects/{w.sid}/cos/123/drilldown"):
        assert client.get(path, headers=w.fh).status_code == 404, path


def test_student_events_are_filtered_by_section(client, ctx, db):
    w = _world(ctx, db)
    a, b = w.students  # sections A and B
    for section in ("A", "B", None):
        r = client.post("/events", headers=w.fh, json={
            "subject_id": w.sid, "type": "quiz", "title": f"for {section}", "section": section,
            "starts_at": "2030-01-01T10:00:00Z"})
        assert r.status_code == 201, r.text
    titles = lambda s: sorted(e["title"] for e in client.get(
        "/events", headers=_tok(ctx, usn=s["usn"], pin=s["pin"])).json())
    assert titles(a) == ["for A", "for None"]
    assert titles(b) == ["for B", "for None"]
    assert len(client.get("/events", headers=w.fh).json()) == 3


# ---------------------------------------------------------------------------
# 2. Upload and status transitions
# ---------------------------------------------------------------------------

def _status(client, w):
    r = client.get(f"/modules/{w.mod['module_id']}/status", headers=w.fh)
    assert r.status_code == 200, r.text
    return r.json()


def _upload(client, w, data=DECK, name="deck.pptx"):
    return client.post(f"/modules/{w.mod['module_id']}/upload", headers=w.fh,
                       files={"file": (name, data, PPTX_MIME)})


def test_upload_pending_to_done_without_ai(client, ctx, db):
    w = _world(ctx, db)
    assert _status(client, w)["processing_status"] == "pending"
    r = _upload(client, w)
    assert r.status_code == 202, r.text
    assert r.json()["processing_status"] == "processing"  # what the client is told straight away
    st = _status(client, w)
    assert st["processing_status"] == "done"
    assert st["chunks_total"] == 3
    # No AI: chunks are kept, untagged, and the upload still succeeds.
    assert st["chunks_tagged"] == 0 and st["chunks_untagged"] == 3
    assert st["error"] is None
    cov = client.get(f"/subjects/{w.sid}/coverage", headers=w.fh).json()
    assert [c["chunk_count"] for c in cov] == [0, 0] and all(c["warning"] for c in cov)


def test_status_is_processing_while_the_job_runs_and_blocks_a_second_upload(client, ctx, db, monkeypatch):
    from backend.app.routers import upload as upload_router
    monkeypatch.setattr(upload_router, "run_upload_pipeline", lambda *a, **k: None)  # job never finishes
    w = _world(ctx, db)
    assert _upload(client, w).status_code == 202
    assert _status(client, w)["processing_status"] == "processing"
    again = _upload(client, w)
    assert again.status_code == 409 and again.json()["error"]["code"] == "CONFLICT"


def test_corrupt_file_ends_failed_with_a_reason(client, ctx, db):
    w = _world(ctx, db)
    assert _upload(client, w, data=b"PK" + b"\x00" * 200).status_code == 202
    st = _status(client, w)
    assert st["processing_status"] == "failed"
    assert st["error"]
    assert st["chunks_total"] == 0
    # A good file afterwards recovers the module.
    assert _upload(client, w).status_code == 202
    assert _status(client, w)["processing_status"] == "done"


def test_upload_input_validation(client, ctx, db):
    w = _world(ctx, db)
    assert _upload(client, w, name="notes.pdf").status_code == 400
    assert _upload(client, w, data=b"not a zip at all").status_code == 400
    assert _status(client, w)["processing_status"] == "pending"  # nothing started


def test_tagging_failure_does_not_fail_the_upload(client, ctx, db, monkeypatch):
    from backend.app.routers import upload as upload_router

    def boom(chunks, cos):
        raise RuntimeError("provider down")

    monkeypatch.setattr(upload_router, "make_tagger", lambda client_, rec: boom)
    w = _world(ctx, db)
    assert _upload(client, w).status_code == 202
    st = _status(client, w)
    assert st["processing_status"] == "done"
    assert st["chunks_total"] == 3 and st["chunks_untagged"] == 3


def test_tags_are_stored_and_coverage_counts_them(client, ctx, db, monkeypatch):
    from backend.app.routers import upload as upload_router

    def tagger(chunks, cos):
        # slide 1,2 -> CO1 ; slide 3 -> CO2 ; slide 1 also CO2
        out = []
        for c in chunks:
            numbers = {1: (1, 2), 2: (1,), 3: (2,)}[c.slide_number]
            out.append({"chunk_id": c.id, "cos": list(numbers), "confidence": 0.9})
        return out

    monkeypatch.setattr(upload_router, "make_tagger", lambda client_, rec: tagger)
    w = _world(ctx, db)
    assert _upload(client, w).status_code == 202
    st = _status(client, w)
    assert st["processing_status"] == "done" and st["chunks_tagged"] == 3
    cov = {c["co_number"]: c for c in client.get(f"/subjects/{w.sid}/coverage", headers=w.fh).json()}
    assert cov[1]["chunk_count"] == 2 and cov[2]["chunk_count"] == 2
    assert not cov[1]["warning"] and not cov[2]["warning"]


def test_reupload_replaces_chunks_instead_of_adding(client, ctx, db):
    w = _world(ctx, db)
    _upload(client, w)
    _upload(client, w, data=_make_deck([("Only slide", "A single slide about stacks and queues.")]))
    st = _status(client, w)
    assert st["processing_status"] == "done" and st["chunks_total"] == 1


# ---------------------------------------------------------------------------
# 3. Quiz submission -> CO sheet consistency
# ---------------------------------------------------------------------------

def _counted_rows(db, w, usn):
    return db.execute(text(
        "SELECT c.number, s.percentage, s.counted, s.attempt_id FROM student_co_scores s "
        "JOIN students st ON st.id = s.student_id JOIN course_outcomes c ON c.id = s.co_id "
        "WHERE st.usn = :u AND s.subject_id = :sid ORDER BY c.number, s.counted DESC"),
        {"u": usn, "sid": w.sid}).fetchall()


def test_sheet_matches_the_counted_flag_after_submissions(client, ctx, db):
    w = _world(ctx, db)
    a, b = w.students
    _attempt(client, ctx, a, w.quiz, ALL_WRONG)          # A attempt 1: 0 / 0
    _attempt(client, ctx, a, w.quiz, ALL_RIGHT)          # A attempt 2: 100 / 100  (best)
    _attempt(client, ctx, b, w.quiz, CO1_RIGHT_CO2_WRONG)  # B attempt 1: 100 / 0

    row_a, _ = _sheet_row(client, w, a["usn"])
    row_b, _ = _sheet_row(client, w, b["usn"])
    assert (row_a["co_scores"]["CO1"], row_a["co_scores"]["CO2"]) == (100.0, 100.0)
    assert (row_b["co_scores"]["CO1"], row_b["co_scores"]["CO2"]) == (100.0, 0.0)

    # Exactly one counted attempt per student and CO, and it equals the sheet.
    for st, row in ((a, row_a), (b, row_b)):
        rows = _counted_rows(db, w, st["usn"])
        for co in (1, 2):
            counted = [r for r in rows if r[0] == co and r[2]]
            assert len(counted) == 1, (st["usn"], co)
            assert float(counted[0][1]) == pytest.approx(row["co_scores"][f"CO{co}"])


def test_changing_one_answer_moves_the_sheet_summary_and_counted_flag(client, ctx, db):
    w = _world(ctx, db)
    a, b = w.students
    _attempt(client, ctx, a, w.quiz, ALL_RIGHT)
    first_attempt, _ = _attempt(client, ctx, b, w.quiz, CO1_RIGHT_CO2_WRONG)

    summary = {s["co_number"]: s for s in client.get(f"/subjects/{w.sid}/co-summary", headers=w.fh).json()}
    assert summary[2]["average"] == pytest.approx(50.0)
    assert (summary[2]["above"], summary[2]["below"]) == (1, 1)
    assert summary[2]["attainment_pct"] == pytest.approx(50.0)
    assert summary[2]["met"] is False  # 50% of students above 60% vs a 60% target
    weak = client.get(f"/subjects/{w.sid}/weak-cos", headers=w.fh).json()
    assert [x["co_number"] for x in weak] == [2]

    # B fixes the two CO2 answers on the second attempt.
    second_attempt, body = _attempt(client, ctx, b, w.quiz, ALL_RIGHT)
    assert body["total_score"] == 4.0

    row_b, _ = _sheet_row(client, w, b["usn"])
    assert row_b["co_scores"]["CO2"] == 100.0
    summary = {s["co_number"]: s for s in client.get(f"/subjects/{w.sid}/co-summary", headers=w.fh).json()}
    assert summary[2]["average"] == pytest.approx(100.0)
    assert summary[2]["met"] is True
    assert client.get(f"/subjects/{w.sid}/weak-cos", headers=w.fh).json() == []
    # The counted flag moved to the new attempt (best rule), still exactly one per CO.
    co2 = [r for r in _counted_rows(db, w, b["usn"]) if r[0] == 2]
    assert [str(r[3]) for r in co2 if r[2]] == [second_attempt]
    assert len([r for r in co2 if r[2]]) == 1


def test_a_third_attempt_is_refused_and_leaves_the_sheet_alone(client, ctx, db):
    w = _world(ctx, db)
    st = w.students[0]
    _attempt(client, ctx, st, w.quiz, ALL_RIGHT)
    _attempt(client, ctx, st, w.quiz, ALL_WRONG)
    h = _tok(ctx, usn=st["usn"], pin=st["pin"])
    r = client.post(f"/quizzes/{w.quiz}/attempts", headers=h)
    assert r.status_code == 409, r.text
    row, _ = _sheet_row(client, w, st["usn"])
    assert row["co_scores"]["CO1"] == 100.0  # best of the two


def test_unattempted_students_show_no_score_not_zero(client, ctx, db):
    w = _world(ctx, db)
    _attempt(client, ctx, w.students[0], w.quiz, ALL_RIGHT)
    row, _ = _sheet_row(client, w, w.students[1]["usn"])
    assert row["co_scores"]["CO1"] is None and row["co_scores"]["CO2"] is None


def test_section_filter_and_module_scope(client, ctx, db):
    w = _world(ctx, db)
    a, b = w.students
    _attempt(client, ctx, a, w.quiz, ALL_RIGHT)
    _attempt(client, ctx, b, w.quiz, ALL_WRONG)
    r = client.get(f"/subjects/{w.sid}/co-sheet", headers=w.fh, params={"section": "A"})
    assert [x["usn"] for x in r.json()["rows"]] == [a["usn"]]
    # scope=module needs a module, and the module must belong to this subject.
    assert client.get(f"/subjects/{w.sid}/co-sheet", headers=w.fh, params={"scope": "module"}).status_code == 400
    ok = client.get(f"/subjects/{w.sid}/co-sheet", headers=w.fh,
                    params={"scope": "module", "module_id": w.mod["module_id"]})
    assert ok.status_code == 200
    foreign = client.get(f"/subjects/{w.sid}/co-sheet", headers=w.fh,
                         params={"scope": "module", "module_id": str(uuid.uuid4())})
    assert foreign.status_code == 404
    assert client.get(f"/subjects/{w.sid}/co-sheet", headers=w.fh, params={"scope": "week"}).status_code == 422


def test_drilldown_shows_the_students_and_questions_behind_the_number(client, ctx, db):
    w = _world(ctx, db)
    a, b = w.students
    _attempt(client, ctx, a, w.quiz, ALL_RIGHT)
    _attempt(client, ctx, b, w.quiz, CO1_RIGHT_CO2_WRONG)
    r = client.get(f"/subjects/{w.sid}/cos/{w.co2['co_id']}/drilldown", headers=w.fh)
    assert r.status_code == 200, r.text
    blob = r.json()
    flat = str(blob)
    assert a["usn"] in flat and b["usn"] in flat
    assert w.qs[2]["question_id"] in flat and w.qs[3]["question_id"] in flat
    assert w.qs[0]["question_id"] not in flat  # CO1 questions are not behind a CO2 number


def test_excel_export_matches_the_on_screen_sheet(client, ctx, db):
    w = _world(ctx, db)
    a, b = w.students
    _attempt(client, ctx, a, w.quiz, ALL_RIGHT)
    _attempt(client, ctx, b, w.quiz, CO1_RIGHT_CO2_WRONG)
    _, sheet = _sheet_row(client, w, a["usn"])

    r = client.get(f"/subjects/{w.sid}/co-sheet/export", headers=w.fh)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in r.headers["content-disposition"]
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    header_idx = next(i for i, row in enumerate(rows) if "USN" in row)
    header = [str(x) for x in rows[header_idx]]
    col = lambda co: next(i for i, h in enumerate(header) if h.startswith(co + " "))  # "CO1 (%)"
    for expected in sheet["rows"]:
        line = next(row for row in rows[header_idx + 1:] if expected["usn"] in row)
        for co, value in expected["co_scores"].items():
            cell = line[col(co)]
            if value is None:
                assert cell in (None, "", "-")
            else:
                assert float(cell) == pytest.approx(value, abs=0.01), (expected["usn"], co)


# ---------------------------------------------------------------------------
# 4. Slide Q&A, study material and study plan
# ---------------------------------------------------------------------------

def _load_chunks(ctx, db, w, tagged=True):
    """Put the DECK slides in the module, tagged to COs, without going through upload."""
    slides = [
        (1, "Binary search tree: a binary search tree keeps smaller keys in the left subtree and larger keys in the right subtree.", 1),
        (2, "Tree traversal: inorder traversal visits left subtree, node, then right subtree and returns keys in sorted order.", 1),
        (3, "Hash tables: a hash table maps keys to buckets using a hash function and resolves collisions by chaining.", 2),
    ]
    for n, body, co in slides:
        cid = db.execute(text(
            "INSERT INTO content_chunks (module_id, slide_number, text, tag_status) "
            "VALUES (:m, :n, :t, :s) RETURNING id"),
            {"m": w.mod["module_id"], "n": n, "t": body, "s": "tagged" if tagged else "untagged"}).scalar()
        if tagged:
            co_id = (w.co1 if co == 1 else w.co2)["co_id"]
            db.execute(text("INSERT INTO chunk_cos (chunk_id, co_id, confidence) VALUES (:c, :o, 0.9)"),
                       {"c": cid, "o": co_id})
    db.commit()


def test_qa_cites_slides_and_refuses_when_nothing_matches(client, ctx, db):
    w = _world(ctx, db)
    _load_chunks(ctx, db, w)
    sh = _tok(ctx, usn=w.students[0]["usn"], pin=w.students[0]["pin"])
    hit = client.post("/qa", headers=sh, json={"subject_id": w.sid, "question": "How does inorder traversal work?"})
    assert hit.status_code == 200, hit.text
    body = hit.json()
    assert body["refused"] is False and body["citations"]
    assert 2 in [c["slide_number"] for c in body["citations"]]

    miss = client.post("/qa", headers=sh, json={"subject_id": w.sid, "question": "Who won the football world cup?"})
    assert miss.status_code == 200
    assert miss.json()["refused"] is True and miss.json()["citations"] == []

    chunks = client.get("/retrieve", headers=sh, params={"subject_id": w.sid, "q": "hash function collisions"}).json()
    assert chunks and chunks[0]["slide_number"] == 3
    assert client.post("/qa", headers=sh, json={"subject_id": w.sid, "question": "ab"}).status_code == 422


def test_study_material_is_generated_once_per_module(client, ctx, db):
    w = _world(ctx, db)
    mid = w.mod["module_id"]
    # No slides yet -> a clear 409, not an empty success.
    assert client.post(f"/modules/{mid}/study-material", headers=w.fh).status_code == 409
    _load_chunks(ctx, db, w)
    first = client.post(f"/modules/{mid}/study-material", headers=w.fh)
    assert first.status_code == 200, first.text
    n_cards = len(first.json()["flashcards"])
    assert n_cards >= 1
    second = client.post(f"/modules/{mid}/study-material", headers=w.fh)
    assert second.status_code == 200
    count = db.execute(text("SELECT count(*) FROM flashcards WHERE module_id = :m"), {"m": mid}).scalar()
    assert count == n_cards  # not duplicated
    # Students of the subject can read it; it is per module, not per student.
    sh = _tok(ctx, usn=w.students[0]["usn"], pin=w.students[0]["pin"])
    assert len(client.get(f"/modules/{mid}/flashcards", headers=sh).json()) == n_cards
    assert client.get(f"/modules/{mid}/notes", headers=sh).status_code == 200
    regen = client.post(f"/modules/{mid}/study-material", headers=w.fh, params={"regenerate": "true"})
    assert regen.status_code == 200
    assert db.execute(text("SELECT count(*) FROM flashcards WHERE module_id = :m"), {"m": mid}).scalar() == n_cards


def test_study_plan_targets_the_students_weak_co_with_real_slides(client, ctx, db):
    w = _world(ctx, db)
    _load_chunks(ctx, db, w)
    client.post(f"/modules/{w.mod['module_id']}/study-material", headers=w.fh)
    a, b = w.students
    _attempt(client, ctx, a, w.quiz, ALL_RIGHT)
    _attempt(client, ctx, b, w.quiz, CO1_RIGHT_CO2_WRONG)   # weak in CO2
    bh = _tok(ctx, usn=b["usn"], pin=b["pin"])
    ah = _tok(ctx, usn=a["usn"], pin=a["pin"])

    assert client.get("/students/me/study-plan", headers=bh, params={"subject_id": w.sid}).status_code == 404
    created = client.post("/students/me/study-plan", headers=bh, json={"subject_id": w.sid})
    assert created.status_code == 201, created.text
    plan = created.json()
    steps = plan["plan"] if isinstance(plan.get("plan"), list) else plan["plan"]["steps"]
    assert [s["co_number"] for s in steps] == [2]
    assert [s["slide_number"] for s in steps[0]["slides"]] == [3]  # only the CO2 slide
    got = client.get("/students/me/study-plan", headers=bh, params={"subject_id": w.sid})
    assert got.status_code == 200 and got.json()["id"] == plan["id"]
    # The student who is not weak has no plan, and cannot see B's.
    assert client.get("/students/me/study-plan", headers=ah, params={"subject_id": w.sid}).status_code == 404


# ---------------------------------------------------------------------------
# 5. Calendar
# ---------------------------------------------------------------------------

def test_calendar_create_list_move_delete_and_validation(client, ctx, db):
    w = _world(ctx, db)
    r = client.post("/events", headers=w.fh, json={
        "subject_id": w.sid, "type": "assignment", "title": "  Lab 1 due  ",
        "starts_at": "2030-02-01T09:00:00Z", "module_id": w.mod["module_id"]})
    assert r.status_code == 201, r.text
    ev = r.json()
    assert ev["title"] == "Lab 1 due"
    moved = client.patch(f"/events/{ev['id']}", headers=w.fh, json={"starts_at": "2030-02-03T09:00:00Z"})
    assert moved.status_code == 200 and moved.json()["starts_at"].startswith("2030-02-03")
    window = client.get("/events", headers=w.fh, params={"from": "2030-02-02T00:00:00Z", "to": "2030-02-04T00:00:00Z"})
    assert [e["id"] for e in window.json()] == [ev["id"]]
    assert client.post("/events", headers=w.fh, json={
        "subject_id": w.sid, "type": "nonsense", "title": "x",
        "starts_at": "2030-02-01T09:00:00Z"}).status_code == 422
    # A module from some other subject cannot be linked.
    other = ctx.make_subject(w.fac["user_id"])
    other_mod = ctx.make_module(other["subject_id"], 1, "Other")
    assert client.post("/events", headers=w.fh, json={
        "subject_id": w.sid, "type": "assignment", "title": "x", "module_id": other_mod["module_id"],
        "starts_at": "2030-02-01T09:00:00Z"}).status_code == 404
    assert client.delete(f"/events/{ev['id']}", headers=w.fh).json() == {"deleted": True}
    assert client.patch(f"/events/{ev['id']}", headers=w.fh, json={"title": "gone"}).status_code == 404


def test_another_faculty_cannot_move_or_delete_my_event(client, ctx, db):
    w = _world(ctx, db)
    ev = client.post("/events", headers=w.fh, json={
        "subject_id": w.sid, "type": "quiz", "title": "Mine", "starts_at": "2030-03-01T09:00:00Z"}).json()
    other = ctx.make_faculty()
    oh = _tok(ctx, email=other["email"], password=other["password"])
    assert client.patch(f"/events/{ev['id']}", headers=oh, json={"title": "hijack"}).status_code == 404
    assert client.delete(f"/events/{ev['id']}", headers=oh).status_code == 404
    still = client.get("/events", headers=w.fh).json()
    assert [e["title"] for e in still] == ["Mine"]


def test_publishing_a_quiz_creates_its_event(client, ctx, db):
    """Aparna's publish endpoint owns this; checked here because the calendar depends on it."""
    w = _world(ctx, db)
    # The world quiz was published straight in SQL, so make a second one through the API.
    now = datetime.now(timezone.utc)
    q = client.post("/quizzes", headers=w.fh, json={
        "module_id": w.mod["module_id"], "title": "API quiz",
        "opens_at": (now + timedelta(days=1)).isoformat(), "closes_at": (now + timedelta(days=2)).isoformat()})
    assert q.status_code in (200, 201), q.text
    qid = q.json().get("id") or q.json().get("quiz_id")
    add = client.post(f"/quizzes/{qid}/questions", headers=w.fh,
                      json={"question_id": w.qs[0]["question_id"]})
    assert add.status_code in (200, 201), add.text
    pub = client.post(f"/quizzes/{qid}/publish", headers=w.fh)
    assert pub.status_code == 200, pub.text
    titles = [e["title"] for e in client.get("/events", headers=w.fh).json()]
    assert any("API quiz" in t for t in titles)
    db.execute(text("DELETE FROM quizzes WHERE id = :q"), {"q": qid})
    db.commit()


# ---------------------------------------------------------------------------
# 6. Usage logging (fake AI client, no credentials)
# ---------------------------------------------------------------------------

class _FakeClient:
    def chat(self, model, messages, temperature=0.2, max_tokens=500, **_):
        return SimpleNamespace(content="Inorder traversal visits left, node, right [Slide 2].",
                               usage={"prompt_tokens": 120, "completion_tokens": 30})

    def close(self):
        pass


def test_ai_calls_are_logged_and_summed_per_step(client, ctx, db, monkeypatch, log_cleanup):
    from backend.app.routers import study as study_router
    from backend.app.services import ai_glue
    monkeypatch.setenv("MODEL_REASONING", "fake-model")
    monkeypatch.setattr(study_router, "get_client", lambda: _FakeClient())
    monkeypatch.setattr(ai_glue, "get_client", lambda: _FakeClient())

    w = _world(ctx, db)
    _load_chunks(ctx, db, w)
    sh = _tok(ctx, usn=w.students[0]["usn"], pin=w.students[0]["pin"])
    before = {r["step"]: r for r in client.get("/usage/summary", headers=w.fh).json()}.get("qa", {})
    for _ in range(2):
        r = client.post("/qa", headers=sh, json={"subject_id": w.sid, "question": "How does inorder traversal work?"})
        assert r.status_code == 200 and r.json()["refused"] is False

    after = {r["step"]: r for r in client.get("/usage/summary", headers=w.fh).json()}["qa"]
    assert after["calls"] - before.get("calls", 0) == 2
    assert after["prompt_tokens"] - before.get("prompt_tokens", 0) == 240
    assert after["completion_tokens"] - before.get("completion_tokens", 0) == 60
    assert after["total_tokens"] - before.get("total_tokens", 0) == 300
    row = db.execute(text("SELECT role, model, step FROM agent_logs WHERE model = 'fake-model' LIMIT 1")).fetchone()
    assert tuple(row) == ("student", "fake-model", "qa")


def test_a_failing_ai_call_falls_back_to_the_slides_and_is_logged(client, ctx, db, monkeypatch, log_cleanup):
    from backend.app.routers import study as study_router

    class Broken(_FakeClient):
        def chat(self, *a, **k):
            raise RuntimeError("provider down")

    monkeypatch.setenv("MODEL_REASONING", "fake-model")
    monkeypatch.setattr(study_router, "get_client", lambda: Broken())
    w = _world(ctx, db)
    _load_chunks(ctx, db, w)
    sh = _tok(ctx, usn=w.students[0]["usn"], pin=w.students[0]["pin"])
    r = client.post("/qa", headers=sh, json={"subject_id": w.sid, "question": "How does inorder traversal work?"})
    assert r.status_code == 200
    assert r.json()["refused"] is False and r.json()["citations"]
    err = db.execute(text("SELECT error FROM agent_logs WHERE model = 'fake-model' AND error IS NOT NULL")).fetchall()
    assert err
