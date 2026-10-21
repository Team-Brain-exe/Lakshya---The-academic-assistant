"""Excel export of the CO sheet (openpyxl).

The workbook is built from the same COSheet the screen shows, so the two cannot
disagree. Layout:

  Sheet 1 "CO sheet"      USN, Name, [Section], one column per CO (percent), Total (percent),
                          and an Average row at the bottom. A CO with no data is left blank,
                          never 0.
  Sheet 2 "Class summary" optional; one row per CO with threshold, target, average,
                          counts above and below the threshold, attainment and whether the
                          target was met.

Values are rounded to 2 decimals, the same as the screen. The university's own format is
not known yet (docs/decisions.md, item 4), so this is the generic layout until faculty
share a sample.
"""

from __future__ import annotations

import io
from typing import Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..engine.co_math import COSheet, COSummary
from ..engine.grading import meets_threshold, round2

_BELOW_FILL = PatternFill("solid", start_color="FFF8D7DA", end_color="FFF8D7DA")
_HEADER_FILL = PatternFill("solid", start_color="FFE9ECEF", end_color="FFE9ECEF")
_BOLD = Font(bold=True)
_NUMBER_FORMAT = "0.00"
_SUMMARY_TITLE = "Class summary"


def _safe_sheet_title(title: str) -> str:
    """Excel sheet names: at most 31 characters, none of [ ] : * ? / \\, no edge apostrophes."""
    cleaned = "".join("-" if ch in "[]:*?/\\" else ch for ch in title).strip().strip("'").strip()
    return (cleaned or "CO sheet")[:31]


def _text(ws, row: int, col: int, value: str | None, bold: bool = False) -> None:
    """Write text. A value like '=1+1' stays text, so a USN or name can never run as a formula."""
    cell = ws.cell(row=row, column=col)
    if value is None or value == "":
        return
    cell.value = value
    cell.data_type = "s"
    if bold:
        cell.font = _BOLD


def _number(ws, row: int, col: int, value: float | None, below: bool = False) -> None:
    """Write a rounded number, or leave the cell blank for 'no data'."""
    if value is None:
        return
    cell = ws.cell(row=row, column=col)
    cell.value = round2(value)
    cell.number_format = _NUMBER_FORMAT
    if below:
        cell.fill = _BELOW_FILL


def _header(ws, labels: Sequence[str]) -> None:
    for col, label in enumerate(labels, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.font = _BOLD
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_sheet_workbook(
    sheet: COSheet,
    thresholds: Mapping[int, float] | None = None,
    summary: Sequence[COSummary] | None = None,
    title: str = "CO sheet",
) -> bytes:
    """Return the .xlsx file as bytes.

    thresholds  CO number -> threshold percent. When given, scores below it are shaded.
    summary     class_summary() output. When given, a second sheet is added.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = _safe_sheet_title(title)

    has_section = any(r.section for r in sheet.rows)
    first_co_col = 4 if has_section else 3
    total_col = first_co_col + len(sheet.co_numbers)

    labels = ["USN", "Name"] + (["Section"] if has_section else [])
    labels += [f"CO{co} (%)" for co in sheet.co_numbers] + ["Total (%)"]
    _header(ws, labels)

    for row_idx, row in enumerate(sheet.rows, start=2):
        _text(ws, row_idx, 1, row.usn)
        _text(ws, row_idx, 2, row.name)
        if has_section:
            _text(ws, row_idx, 3, row.section)
        for offset, co in enumerate(sheet.co_numbers):
            value = row.co_scores[co]
            below = (
                thresholds is not None
                and co in thresholds
                and value is not None
                and not meets_threshold(value, thresholds[co])
            )
            _number(ws, row_idx, first_co_col + offset, value, below)
        _number(ws, row_idx, total_col, row.total_pct)

    avg_row = len(sheet.rows) + 2
    _text(ws, avg_row, 1, "Average", bold=True)
    for offset, co in enumerate(sheet.co_numbers):
        _number(ws, avg_row, first_co_col + offset, sheet.averages[co])
        ws.cell(row=avg_row, column=first_co_col + offset).font = _BOLD

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 24
    for col in range(3, total_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 12
    ws.freeze_panes = ws.cell(row=2, column=3)

    if summary is not None:
        ws2 = wb.create_sheet(_SUMMARY_TITLE)
        _header(
            ws2,
            [
                "CO",
                "Threshold (%)",
                "Target (%)",
                "Enrolled",
                "Attempted",
                "Average (%)",
                "At or above threshold",
                "Below threshold",
                "Attainment (%)",
                "Target met",
            ],
        )
        for row_idx, s in enumerate(summary, start=2):
            _text(ws2, row_idx, 1, f"CO{s.co}")
            _number(ws2, row_idx, 2, s.threshold_pct)
            _number(ws2, row_idx, 3, s.target_pct)
            ws2.cell(row=row_idx, column=4, value=s.enrolled)
            ws2.cell(row=row_idx, column=5, value=s.attempted)
            _number(ws2, row_idx, 6, s.average)
            ws2.cell(row=row_idx, column=7, value=s.above)
            ws2.cell(row=row_idx, column=8, value=s.below)
            _number(ws2, row_idx, 9, s.attainment_pct)
            verdict = "No data" if not s.has_data else ("Yes" if s.met else "No")
            _text(ws2, row_idx, 10, verdict)
        for col in range(1, 11):
            ws2.column_dimensions[get_column_letter(col)].width = 16
        ws2.freeze_panes = ws2.cell(row=2, column=2)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
