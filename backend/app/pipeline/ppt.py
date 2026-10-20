"""Split an uploaded PowerPoint file into content chunks, one per slide.

Each chunk is the slide number plus the slide's text: the title first, then text
boxes, table cells (one row per line, cells joined with " | "), text inside groups,
and the speaker notes. Slides with no text (for example picture-only slides) are
skipped and reported, so the caller can tell faculty.
"""

from __future__ import annotations

import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import BinaryIO

from pptx import Presentation
from pptx.exc import PackageNotFoundError
from pptx.shapes.group import GroupShape


class PptParseError(Exception):
    """The uploaded file could not be read as a .pptx presentation."""


@dataclass(frozen=True)
class SlideChunk:
    slide_number: int  # 1-based, as shown in PowerPoint
    text: str


@dataclass(frozen=True)
class ParsedDeck:
    chunks: tuple[SlideChunk, ...]
    skipped_slides: tuple[int, ...]  # slide numbers with no text
    total_slides: int


def _clean_lines(raw: str) -> list[str]:
    """Split on newlines and PowerPoint's vertical-tab line breaks; drop blanks."""
    lines = []
    for line in raw.replace("\x0b", "\n").replace("\r", "\n").split("\n"):
        line = " ".join(line.split())
        if line:
            lines.append(line)
    return lines


def _shape_lines(shape) -> list[str]:
    """Text lines from one shape, going into groups and tables."""
    if isinstance(shape, GroupShape):
        lines: list[str] = []
        for child in shape.shapes:
            lines.extend(_shape_lines(child))
        return lines

    lines = []
    if getattr(shape, "has_text_frame", False) and shape.text_frame is not None:
        lines.extend(_clean_lines(shape.text_frame.text))
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [" ".join(c.text.split()) for c in row.cells]
            cells = [c for c in cells if c]
            if cells:
                lines.append(" | ".join(cells))
    return lines


def _slide_text(slide, include_notes: bool) -> str:
    lines: list[str] = []

    title = slide.shapes.title
    title_id = None
    if title is not None and getattr(title, "has_text_frame", False):
        lines.extend(_clean_lines(title.text_frame.text))
        title_id = title.shape_id

    for shape in slide.shapes:
        if title_id is not None and shape.shape_id == title_id:
            continue
        lines.extend(_shape_lines(shape))

    if include_notes and slide.has_notes_slide:
        frame = slide.notes_slide.notes_text_frame
        if frame is not None:
            notes = _clean_lines(frame.text)
            if notes:
                lines.append("Speaker notes: " + " ".join(notes))

    return "\n".join(lines)


def parse_pptx(source: str | Path | BinaryIO, include_notes: bool = True) -> ParsedDeck:
    """Read a .pptx and return one chunk per slide that has text.

    source: a file path or an open binary file (for example an uploaded file).
    Raises PptParseError if the file is not a valid .pptx.
    """
    try:
        prs = Presentation(source)
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise PptParseError(f"Could not read the file as a .pptx presentation: {exc}") from exc

    chunks: list[SlideChunk] = []
    skipped: list[int] = []
    total = 0
    for number, slide in enumerate(prs.slides, start=1):
        total = number
        text = _slide_text(slide, include_notes)
        if text.strip():
            chunks.append(SlideChunk(slide_number=number, text=text))
        else:
            skipped.append(number)
    return ParsedDeck(chunks=tuple(chunks), skipped_slides=tuple(skipped), total_slides=total)
